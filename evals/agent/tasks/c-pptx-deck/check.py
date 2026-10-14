"""Pass: three slides titled Agenda, Results, Next Steps; the Results body has exactly the three
points. A slide's title is its title placeholder, or else its topmost text shape (libraries such as
pptxgenjs write titles as text boxes)."""
import os, sys
from pptx import Presentation

p = os.path.join(sys.argv[1], "deck.pptx")
if not os.path.exists(p):
    sys.exit("no deck.pptx")
pr = Presentation(p)


def title_and_body(slide):
    texts = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = slide.shapes.title if slide.shapes.title is not None and slide.shapes.title.text.strip() else (
        min(texts, key=lambda sh: (sh.top or 0, sh.left or 0)) if texts else None)
    # python-pptx builds a new proxy per access: compare the underlying XML elements
    tid = title._element if title is not None else None
    body = [pa.text.strip() for sh in texts if sh._element is not tid for pa in sh.text_frame.paragraphs if pa.text.strip()]
    return (title.text.strip() if title is not None else ""), body


slides = [title_and_body(s) for s in pr.slides]
if [t for t, _ in slides] != ["Agenda", "Results", "Next Steps"]:
    sys.exit(f"titles {[t for t, _ in slides]}")
if slides[1][1] != ["Revenue up 12%", "Churn down 3%", "NPS 41"]:
    sys.exit(f"Results body {slides[1][1]}")
print("ok")

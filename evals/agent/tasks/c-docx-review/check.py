"""Pass: heading style, the 3x3 table, and a two-item bulleted list in a list style."""
import os, sys
from docx import Document

p = os.path.join(sys.argv[1], "report.docx")
if not os.path.exists(p):
    sys.exit("no report.docx")
d = Document(p)
heads = [x for x in d.paragraphs if x.text.strip() == "Quarterly Review"]
if not heads or not (heads[0].style.name.startswith("Heading") or heads[0].style.name == "Title"):
    sys.exit("heading missing or not a heading style")
want = [["Region", "Q1", "Q2"], ["North", "10", "12"], ["South", "8", "9"]]
if not any([[c.text.strip() for c in r.cells] for r in t.rows] == want for t in d.tables):
    sys.exit("table missing or different")
items = [x for x in d.paragraphs if x.text.strip() in ("Hire two engineers", "Ship the mobile app")]
bullet = [x for x in items if "List" in x.style.name or x._p.pPr is not None and x._p.pPr.numPr is not None]
if len(bullet) != 2:
    sys.exit(f"bulleted items found: {len(bullet)}")
print("ok")

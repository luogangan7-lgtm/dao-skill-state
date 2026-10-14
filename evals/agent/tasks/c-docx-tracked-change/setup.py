"""contract.docx: a heading and three paragraphs; the second holds the payment term."""
from docx import Document

d = Document()
d.add_heading("Services Agreement", level=1)
d.add_paragraph("The Supplier provides the services described in Schedule A.")
d.add_paragraph("Invoices are payable within 30 days of receipt.")
d.add_paragraph("Either party may end this agreement with written notice.")
d.save("contract.docx")

"""Pass: the workbook has the inputs and live formulas whose values are right."""
import os, re, sys
from openpyxl import load_workbook

p = os.path.join(sys.argv[1], "budget.xlsx")
if not os.path.exists(p):
    sys.exit("no budget.xlsx")
wb = load_workbook(p)
if "Q1" not in wb.sheetnames:
    sys.exit("no sheet Q1")
ws = wb["Q1"]


def val(ref, seen=()):
    v = ws[ref].value
    if isinstance(v, str) and v.startswith("="):
        if ref in seen:
            raise ValueError("cycle")
        return ev(v[1:], seen + (ref,))
    return v


def ev(expr, seen):
    def rng(m):
        (c1, r1), (c2, r2) = (re.match(r"([A-Z]+)(\d+)", g.replace("$", "")).groups() for g in m.groups())
        cells = [f"{chr(c)}{r}" for c in range(ord(c1), ord(c2) + 1) for r in range(int(r1), int(r2) + 1)]
        return "(" + "+".join(str(val(c, seen)) for c in cells) + ")"
    e = re.sub(r"(?i)SUM\(\s*(\$?[A-Z]+\$?\d+)\s*:\s*(\$?[A-Z]+\$?\d+)\s*\)", rng, expr)
    e = re.sub(r"\$?([A-Z]+)\$?(\d+)", lambda m: str(val(m.group(1) + m.group(2), seen)), e)
    if not re.fullmatch(r"[0-9+\-*/(). ]+", e):
        raise ValueError(f"unsupported formula {expr!r}")
    return eval(e)


want = {"B2": 1200, "C2": 1350, "D2": 1500, "B3": 800, "C3": 820, "D3": 900}
for ref, v in want.items():
    if val(ref) != v:
        sys.exit(f"{ref} = {ws[ref].value!r}, want {v}")
expect = {"B4": 400, "C4": 530, "D4": 600, "E2": 4050, "E3": 2520, "E4": 1530}
for ref, v in expect.items():
    raw = ws[ref].value
    if not (isinstance(raw, str) and raw.startswith("=")):
        sys.exit(f"{ref} is not a formula: {raw!r}")
    try:
        got = val(ref)
    except ValueError as e:
        sys.exit(f"{ref}: {e}")
    if abs(got - v) > 1e-9:
        sys.exit(f"{ref} evaluates to {got}, want {v}")
print("ok")
